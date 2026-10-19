#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""FileForge — boîte à outils fichiers pour macOS.
PDF, images, vidéo, audio, documents, tableurs, archives.
Usage : fileforge.py fichier1 [fichier2 ...]  (affiche un menu natif)
"""
import os, sys, re, io, json, math, shutil, subprocess, traceback, datetime, zipfile, tarfile, gzip, tempfile
from pathlib import Path

APP = "FileForge"
__version__ = "1.0.0"
HOME_DIR = Path(__file__).resolve().parent
IS_WIN = sys.platform == "win32"
IS_MAC = sys.platform == "darwin"
IS_LINUX = sys.platform.startswith("linux")
if IS_WIN:
    LOG = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "FileForge" / "FileForge.log"
elif IS_LINUX:
    LOG = Path(os.environ.get("XDG_STATE_HOME") or Path.home() / ".local/state") / "fileforge" / "FileForge.log"
else:
    LOG = Path.home() / "Library/Logs/FileForge.log"
try:
    LOG.parent.mkdir(parents=True, exist_ok=True)
except Exception:
    pass
# langues OCR fournies par l'installateur (toutes plateformes)
if (HOME_DIR / "tessdata").is_dir():
    os.environ.setdefault("TESSDATA_PREFIX", str(HOME_DIR / "tessdata"))
if IS_LINUX and "TESSDATA_PREFIX" not in os.environ:
    for _td in ("/usr/share/tessdata", "/usr/share/tesseract-ocr/5/tessdata", "/usr/share/tesseract-ocr/4.00/tessdata"):
        if os.path.isdir(_td):
            os.environ["TESSDATA_PREFIX"] = _td
            break
NOTES = []

if IS_WIN:
    # langues OCR fournies avec FileForge
    if (HOME_DIR / "tessdata").is_dir():
        os.environ.setdefault("TESSDATA_PREFIX", str(HOME_DIR / "tessdata"))
    # aucune fenêtre de console quand on lance ffmpeg, LibreOffice…
    _Popen = subprocess.Popen

    class _HiddenPopen(_Popen):
        def __init__(self, *a, **k):
            k.setdefault("creationflags", 0x08000000)  # CREATE_NO_WINDOW
            super().__init__(*a, **k)
    subprocess.Popen = _HiddenPopen


def open_path(p):
    if IS_WIN:
        os.startfile(str(p))
    elif IS_LINUX:
        subprocess.Popen(["xdg-open", str(p)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                         start_new_session=True)
    else:
        subprocess.run(["open", str(p)])


def reveal(p):
    if IS_WIN:
        subprocess.Popen(f'explorer /select,"{p}"')
    elif IS_LINUX:
        # interface standard des gestionnaires de fichiers (Nautilus, Dolphin, Thunar, Nemo…)
        r = subprocess.run(["dbus-send", "--session", "--print-reply", "--dest=org.freedesktop.FileManager1",
                            "--type=method_call", "/org/freedesktop/FileManager1",
                            "org.freedesktop.FileManager1.ShowItems",
                            f"array:string:{Path(p).resolve().as_uri()}", "string:"],
                           capture_output=True, timeout=5)
        if r.returncode != 0:
            open_path(Path(p).resolve().parent)
    else:
        subprocess.run(["open", "-R", str(p)])


def font_file():
    for c in ("/System/Library/Fonts/Helvetica.ttc", r"C:\Windows\Fonts\segoeui.ttf",
              r"C:\Windows\Fonts\arial.ttf", "/usr/share/fonts/TTF/DejaVuSans.ttf",
              "/usr/share/fonts/noto/NotoSans-Regular.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
              "/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/liberation/LiberationSans-Regular.ttf"):
        if os.path.exists(c):
            return c
    return "DejaVuSans.ttf"

# ───────────────────────── utilitaires UI (AppleScript) ─────────────────────────

class Cancel(Exception):
    pass


def log(msg):
    try:
        with open(LOG, "a") as f:
            f.write(f"[{datetime.datetime.now():%Y-%m-%d %H:%M:%S}] {msg}\n")
    except Exception:
        pass


def _q(s):
    return '"' + str(s).replace("\\", "\\\\").replace('"', '\\"') + '"'


def osa(script):
    r = subprocess.run(["osascript", "-e", script], capture_output=True, text=True)
    if r.returncode != 0:
        if "-128" in r.stderr:
            raise Cancel()
        raise RuntimeError(r.stderr.strip())
    return r.stdout.rstrip("\n")


def choose(items, prompt, default=None, multiple=False, ok="OK"):
    lst = "{" + ",".join(_q(i) for i in items) + "}"
    d = f" default items {{{_q(default if default in items else items[0])}}}" if not multiple else ""
    m = " with multiple selections allowed" if multiple else ""
    res = osa(
        f"activate\nset r to choose from list {lst} with title {_q(APP)} with prompt {_q(prompt)}{d}{m} "
        f'OK button name {_q(ok)} cancel button name "Annuler"\n'
        "if r is false then error number -128\n"
        "set AppleScript's text item delimiters to linefeed\nreturn r as text")
    return res.split("\n") if multiple else res


def ask(prompt, default="", hidden=False):
    h = " with hidden answer" if hidden else ""
    return osa(
        f"activate\nreturn text returned of (display dialog {_q(prompt)} default answer {_q(default)} "
        f'with title {_q(APP)}{h} buttons {{"Annuler","OK"}} default button "OK" cancel button "Annuler")').strip()


def buttons(prompt, btns, default=None):
    b = "{" + ",".join(_q(x) for x in btns) + "}"
    c = ' cancel button "Annuler"' if "Annuler" in btns else ""
    return osa(
        f"activate\nreturn button returned of (display dialog {_q(prompt)} with title {_q(APP)} "
        f"buttons {b} default button {_q(default or btns[-1])}{c})")


def alert(title, msg=""):
    try:
        osa(f"activate\ndisplay alert {_q(title)} message {_q(msg)} as warning")
    except Exception:
        pass


def notify(msg):
    try:
        osa(f"display notification {_q(msg)} with title {_q(APP)}")
    except Exception:
        pass


def choose_file(prompt, types=None):
    t = (" of type {" + ",".join(_q(x) for x in types) + "}") if types else ""
    return osa(f"activate\nreturn POSIX path of (choose file with prompt {_q(prompt)}{t})")

# ───────────────────────── utilitaires fichiers ─────────────────────────

PDF = {"pdf"}
IMG = {"jpg", "jpeg", "png", "heic", "heif", "webp", "tif", "tiff", "bmp", "gif", "avif"}
VID = {"mp4", "mov", "m4v", "mkv", "avi", "webm", "wmv", "flv", "mpg", "mpeg", "3gp"}
AUD = {"mp3", "wav", "m4a", "aac", "flac", "ogg", "opus", "aiff", "aif", "wma"}
WORDS = {"docx", "doc", "rtf", "odt"}
OFFICE = WORDS | {"pptx", "ppt", "odp", "xlsx", "xls", "ods"}
TEXT = {"txt", "md", "markdown", "html", "htm"}
TAB = {"csv", "tsv", "xlsx", "xls", "ods", "json"}
ARC = {"zip", "tar", "gz", "tgz", "bz2", "xz"}


def ext(p):
    p = str(p)
    for double in (".tar.gz", ".tar.bz2", ".tar.xz"):
        if p.lower().endswith(double):
            return "tar"
    return Path(p).suffix.lower().lstrip(".")


def natural(p):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", Path(p).name.lower())]


def unique(p):
    p = Path(p)
    if not p.exists():
        return p
    i = 2
    while True:
        c = p.with_name(f"{p.stem} ({i}){p.suffix}")
        if not c.exists():
            return c
        i += 1


def out(src, tag, new_ext=None):
    src = Path(src)
    e = new_ext or src.suffix.lstrip(".")
    return unique(src.with_name(f"{src.stem}_{tag}.{e}"))


def outdir(src, tag):
    d = unique(Path(src).with_name(f"{Path(src).stem}_{tag}"))
    d.mkdir(parents=True)
    return d


def human(n):
    for u in ("o", "Ko", "Mo", "Go"):
        if n < 1024:
            return f"{n:.0f} {u}" if u == "o" else f"{n:.1f} {u}"
        n /= 1024
    return f"{n:.1f} To"


def size_note(src, dst):
    a, b = Path(src).stat().st_size, Path(dst).stat().st_size
    pct = (1 - b / a) * 100 if a else 0
    if b < a:
        NOTES.append(f"{Path(src).name} : {human(a)} → {human(b)} (−{pct:.0f} %)")
    else:
        NOTES.append(f"{Path(src).name} : pas de gain possible ({human(a)} → {human(b)}).")


def parse_ranges(s, n):
    """'1-3, 5, 8-fin' -> indices 0-based. Supporte les plages décroissantes (5-1)."""
    s = s.replace(" ", "").lower()
    if s in ("", "tout", "toutes", "all", "*"):
        return list(range(n))
    res = []
    for part in s.split(","):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            a = 1 if a == "" else (n if a in ("fin", "end") else int(a))
            b = n if b in ("", "fin", "end") else int(b)
            step = 1 if b >= a else -1
            res += list(range(a, b + step, step))
        elif part in ("fin", "end"):
            res.append(n)
        else:
            res.append(int(part))
    bad = [p for p in res if p < 1 or p > n]
    if bad:
        raise ValueError(f"Pages hors limites : {bad} (le document a {n} pages).")
    return [p - 1 for p in res]


def parse_time(s):
    s = s.strip().replace(",", ".")
    if not s:
        return None
    parts = [float(x) for x in s.split(":")]
    t = 0.0
    for x in parts:
        t = t * 60 + x
    return t


def which(name, *extra):
    p = shutil.which(name)
    if p:
        return p
    for c in extra:
        if os.path.exists(c):
            return c
    return None


FFMPEG = which("ffmpeg", "/opt/homebrew/bin/ffmpeg", "/usr/local/bin/ffmpeg")
if not FFMPEG:
    try:
        import imageio_ffmpeg
        FFMPEG = imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        FFMPEG = None


def ff(*args):
    if not FFMPEG:
        raise RuntimeError("ffmpeg est introuvable." + ("" if IS_WIN else " Installe-le avec : brew install ffmpeg"))
    r = subprocess.run([FFMPEG, "-y", "-hide_banner", "-loglevel", "error", *map(str, args)],
                       capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError("ffmpeg : " + (r.stderr.strip().splitlines() or ["erreur"])[-1])

# ───────────────────────── registre des actions ─────────────────────────

ACTIONS = []


def action(label, types=None, minf=1):
    def deco(fn):
        ACTIONS.append(dict(label=label, types=types, minf=minf, fn=fn))
        return fn
    return deco

# ───────────────────────── fenêtres visuelles (Tk) ─────────────────────────

def _front(root):
    try:
        import AppKit
        AppKit.NSApplication.sharedApplication().activateIgnoringOtherApps_(True)
    except Exception:
        pass
    root.lift()
    root.attributes("-topmost", True)
    root.after(600, lambda: root.attributes("-topmost", False))
    root.focus_force()


def pick_regions(src, kind, title, hint, multi=False, start_page=0):
    """Fenêtre où l'utilisateur dessine des rectangles.
    kind='pdf' (src = document pymupdf) ou 'img' (src = image PIL).
    Retourne {page_index: [(x0,y0,x1,y1), ...]} en coordonnées source."""
    import tkinter as tk
    from PIL import Image, ImageTk
    root = tk.Tk()
    root.title(f"{APP} — {title}")
    sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
    MW, MH = int(sw * 0.8), int(sh * 0.70)
    if kind == "pdf":
        import pymupdf
        n = src.page_count

        def render(i):
            R = src[i].rect
            z = min(MW / R.width, MH / R.height)
            pix = src[i].get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=False)
            return Image.frombytes("RGB", (pix.width, pix.height), pix.samples), z
    else:
        n = 1
        base = src.convert("RGB")

        def render(i):
            z = min(MW / base.width, MH / base.height, 1.0)
            return base.resize((max(1, int(base.width * z)), max(1, int(base.height * z)))), z

    st = {"i": min(start_page, n - 1), "z": 1.0, "tk": None, "w": 1, "h": 1}
    sel = {}
    tk.Label(root, text=hint, justify="left", font=("Helvetica", 13)).pack(padx=14, pady=(12, 6), anchor="w")
    c = tk.Canvas(root, highlightthickness=1, highlightbackground="#999", cursor="crosshair")
    c.pack(padx=14)
    bar = tk.Frame(root)
    bar.pack(fill="x", padx=14, pady=12)
    lbl = tk.Label(bar, font=("Helvetica", 12))

    def show():
        img, z = render(st["i"])
        st.update(z=z, tk=ImageTk.PhotoImage(img), w=img.width, h=img.height)
        c.config(width=img.width, height=img.height)
        c.delete("all")
        c.create_image(0, 0, image=st["tk"], anchor="nw")
        for r in sel.get(st["i"], []):
            c.create_rectangle(*[v * z for v in r], outline="#e11d48", width=3)
        lbl.config(text=f"Page {st['i'] + 1} / {n}")

    drag = {"s": None, "id": None}

    def down(e):
        drag["s"] = (e.x, e.y)
        drag["id"] = c.create_rectangle(e.x, e.y, e.x, e.y, outline="#e11d48", width=2, dash=(6, 3))

    def move(e):
        if drag["id"]:
            x0, y0 = drag["s"]
            c.coords(drag["id"], x0, y0, e.x, e.y)

    def up(e):
        if not drag["id"]:
            return
        x0, y0 = drag["s"]
        drag["id"] = None
        W, H = st["w"], st["h"]
        x0, x1 = sorted((max(0, min(W, x0)), max(0, min(W, e.x))))
        y0, y1 = sorted((max(0, min(H, y0)), max(0, min(H, e.y))))
        if x1 - x0 >= 6 and y1 - y0 >= 6:
            z = st["z"]
            r = (x0 / z, y0 / z, x1 / z, y1 / z)
            if multi:
                sel.setdefault(st["i"], []).append(r)
            else:
                sel.clear()
                sel[st["i"]] = [r]
        show()

    c.bind("<ButtonPress-1>", down)
    c.bind("<B1-Motion>", move)
    c.bind("<ButtonRelease-1>", up)

    def nav(d):
        st["i"] = (st["i"] + d) % n
        show()

    if n > 1:
        tk.Button(bar, text="◀", width=3, command=lambda: nav(-1)).pack(side="left")
        lbl.pack(side="left", padx=8)
        tk.Button(bar, text="▶", width=3, command=lambda: nav(1)).pack(side="left")
        root.bind("<Left>", lambda e: nav(-1))
        root.bind("<Right>", lambda e: nav(1))
    res = {"ok": False}

    def clear():
        if multi:
            sel.pop(st["i"], None)
        else:
            sel.clear()
        show()

    def ok():
        if not sel:
            root.bell()
            return
        res["ok"] = True
        root.destroy()

    tk.Button(bar, text="Valider ✓", command=ok, default="active").pack(side="right")
    tk.Button(bar, text="Annuler", command=root.destroy).pack(side="right", padx=6)
    tk.Button(bar, text="Effacer", command=clear).pack(side="right")
    root.bind("<Return>", lambda e: ok())
    root.bind("<Escape>", lambda e: root.destroy())
    show()
    _front(root)
    root.mainloop()
    if not res["ok"]:
        raise Cancel()
    return sel


SIG_FILE = HOME_DIR / "signature.png"


def draw_signature():
    import tkinter as tk
    from PIL import Image, ImageDraw
    W, H, S = 640, 240, 3
    big = Image.new("RGBA", (W * S, H * S), (0, 0, 0, 0))
    dr = ImageDraw.Draw(big)
    ink = (12, 30, 90, 255)
    root = tk.Tk()
    root.title(f"{APP} — Dessine ta signature")
    tk.Label(root, text="Signe avec la souris ou le trackpad, puis clique « Valider ».",
             font=("Helvetica", 13)).pack(padx=14, pady=(12, 6))
    c = tk.Canvas(root, width=W, height=H, bg="white", highlightthickness=1, highlightbackground="#999", cursor="pencil")
    c.pack(padx=14)
    st = {"last": None, "drawn": False, "ok": False}

    def dot(x, y):
        r = 1.6 * S
        dr.ellipse([x * S - r, y * S - r, x * S + r, y * S + r], fill=ink)

    def down(e):
        st["last"] = (e.x, e.y)
        dot(e.x, e.y)

    def move(e):
        if st["last"]:
            x0, y0 = st["last"]
            c.create_line(x0, y0, e.x, e.y, width=3, capstyle="round", fill="#0c1e5a")
            dr.line([(x0 * S, y0 * S), (e.x * S, e.y * S)], fill=ink, width=int(3.2 * S))
            dot(e.x, e.y)
            st["last"] = (e.x, e.y)
            st["drawn"] = True

    def up(e):
        st["last"] = None

    def clear():
        c.delete("all")
        dr.rectangle([0, 0, W * S, H * S], fill=(0, 0, 0, 0))
        st["drawn"] = False

    def ok():
        if st["drawn"]:
            st["ok"] = True
            root.destroy()

    c.bind("<ButtonPress-1>", down)
    c.bind("<B1-Motion>", move)
    c.bind("<ButtonRelease-1>", up)
    bar = tk.Frame(root)
    bar.pack(fill="x", padx=14, pady=12)
    tk.Button(bar, text="Valider ✓", command=ok).pack(side="right")
    tk.Button(bar, text="Annuler", command=root.destroy).pack(side="right", padx=6)
    tk.Button(bar, text="Effacer", command=clear).pack(side="left")
    _front(root)
    root.mainloop()
    if not st["ok"]:
        raise Cancel()
    img = big.crop(big.getbbox())
    img.save(SIG_FILE)
    return SIG_FILE

# ───────────────────────── PDF ─────────────────────────

def open_pdf(path):
    import pymupdf
    d = pymupdf.open(str(path))
    if d.needs_pass:
        pw = ask(f"« {Path(path).name} » est protégé par un mot de passe.\nMot de passe :", hidden=True)
        if not d.authenticate(pw):
            raise RuntimeError("Mot de passe incorrect.")
    return d


def save_pdf(doc, dst, **kw):
    kw.setdefault("garbage", 3)
    kw.setdefault("deflate", True)
    doc.save(str(dst), **kw)
    return dst


def _heif():
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        try:
            pillow_heif.register_avif_opener()
        except Exception:
            pass
    except Exception:
        pass


def img_jpeg_bytes(f):
    from PIL import Image, ImageOps
    _heif()
    im = ImageOps.exif_transpose(Image.open(f))
    im = flatten(im)
    b = io.BytesIO()
    im.save(b, "JPEG", quality=92)
    return b.getvalue(), im.size


def add_image_page(doc, f, fit_a4=True):
    import pymupdf
    data, (w, h) = img_jpeg_bytes(f)
    if fit_a4:
        r = pymupdf.paper_rect("a4")
        if w > h:
            r = pymupdf.Rect(0, 0, r.height, r.width)
        page = doc.new_page(width=r.width, height=r.height)
        page.insert_image(page.rect + (28, 28, -28, -28), stream=data, keep_proportion=True)
    else:
        page = doc.new_page(width=w * 0.75, height=h * 0.75)
        page.insert_image(page.rect, stream=data)


def crop_page(page, l, t, r, b):
    import pymupdf
    R = page.rect
    v = pymupdf.Rect(R.x0 + l, R.y0 + t, R.x1 - r, R.y1 - b)
    if v.is_empty or v.width < 10 or v.height < 10:
        raise ValueError("La zone à conserver est trop petite.")
    u = v * page.derotation_matrix
    u.normalize()
    cb = page.cropbox
    new = pymupdf.Rect(cb.x0 + u.x0, cb.y0 + u.y0, cb.x0 + u.x1, cb.y0 + u.y1)
    page.set_cropbox(new & page.mediabox)


def content_bbox(page):
    import pymupdf
    R = page.rect
    area = R.width * R.height
    bb = pymupdf.Rect()
    for kind, rect in page.get_bboxlog():
        r = pymupdf.Rect(rect) & R
        if r.is_empty or r.width * r.height > 0.92 * area:
            continue  # ignore les fonds pleine page
        bb |= r
    return bb


@action("📄 PDF · Fusionner en un seul PDF", PDF | IMG, minf=2)
def pdf_merge(files):
    import pymupdf
    od = pymupdf.open()
    for f in files:
        if ext(f) == "pdf":
            od.insert_pdf(open_pdf(f))
        else:
            add_image_page(od, f)
    NOTES.append("Ordre : alphabétique (renomme les fichiers 01_, 02_… pour imposer un ordre).")
    return [save_pdf(od, unique(Path(files[0]).with_name("Fusion.pdf")))]


@action("📄 PDF · Diviser", PDF)
def pdf_split(files):
    import pymupdf
    mode = choose(["Une page = un PDF", "Toutes les N pages", "Selon des plages (ex : 1-3; 4-8; 9-fin)"],
                  "Comment diviser ?")
    k = int(ask("Nombre de pages par fichier :", "2")) if mode.startswith("Toutes") else 0
    outs = []
    for f in files:
        doc = open_pdf(f)
        n = doc.page_count
        if mode.startswith("Une"):
            groups = [[i] for i in range(n)]
        elif k:
            groups = [list(range(i, min(i + k, n))) for i in range(0, n, k)]
        else:
            s = ask(f"« {Path(f).name} » ({n} pages)\nPlages séparées par des points-virgules :", "1-2; 3-fin")
            groups = [parse_ranges(g, n) for g in s.split(";") if g.strip()]
        d = outdir(f, "divisé")
        for gi, g in enumerate(groups, 1):
            nd = pymupdf.open()
            for p in g:
                nd.insert_pdf(doc, from_page=p, to_page=p)
            lab = f"p{g[0] + 1}" if len(g) == 1 else f"p{g[0] + 1}-{g[-1] + 1}"
            save_pdf(nd, d / f"{Path(f).stem}_{gi:03d}_{lab}.pdf")
        outs.append(d)
    return outs


@action("📄 PDF · Extraire des pages", PDF)
def pdf_extract(files):
    outs = []
    for f in files:
        doc = open_pdf(f)
        sel = parse_ranges(ask(f"Pages à extraire de « {Path(f).name} » ({doc.page_count} pages)\n"
                               "ex : 1-3, 5, 8-fin", "1"), doc.page_count)
        doc.select(sel)
        outs.append(save_pdf(doc, out(f, "extrait")))
    return outs


@action("📄 PDF · Supprimer des pages", PDF)
def pdf_delete(files):
    outs = []
    for f in files:
        doc = open_pdf(f)
        n = doc.page_count
        s = ask(f"Pages à SUPPRIMER de « {Path(f).name} » ({n} pages)\nex : 2, 5-7", "")
        if not s.strip() or s.strip().lower() in ("tout", "toutes", "*"):
            raise ValueError("Indique les numéros des pages à supprimer (ex : 2, 5-7).")
        rm = set(parse_ranges(s, n))
        keep = [i for i in range(n) if i not in rm]
        if not keep:
            raise ValueError("Impossible de supprimer toutes les pages.")
        doc.select(keep)
        outs.append(save_pdf(doc, out(f, "modifié")))
    return outs


@action("📄 PDF · Réorganiser / inverser les pages", PDF)
def pdf_reorder(files):
    mode = choose(["Inverser l'ordre des pages", "Ordre personnalisé…"], "Réorganiser :")
    outs = []
    for f in files:
        doc = open_pdf(f)
        n = doc.page_count
        if mode.startswith("Inverser"):
            order = list(range(n - 1, -1, -1))
        else:
            order = parse_ranges(ask(f"Nouvel ordre pour « {Path(f).name} » ({n} pages)\n"
                                     "ex : 3, 1, 2, 4-fin   (une page peut être répétée)", f"1-{n}"), n)
        doc.select(order)
        outs.append(save_pdf(doc, out(f, "réorganisé")))
    return outs


@action("📄 PDF · Pivoter des pages", PDF)
def pdf_rotate(files):
    a = choose(["90° vers la droite", "90° vers la gauche", "180°"], "Rotation :")
    angle = {"90° vers la droite": 90, "90° vers la gauche": 270, "180°": 180}[a]
    pages = ask("Pages à pivoter (ex : 1, 3-5) — « tout » pour toutes :", "tout")
    outs = []
    for f in files:
        doc = open_pdf(f)
        for p in parse_ranges(pages, doc.page_count):
            pg = doc[p]
            pg.set_rotation((pg.rotation + angle) % 360)
        outs.append(save_pdf(doc, out(f, "pivoté")))
    return outs


@action("📄 PDF · Compresser", PDF)
def pdf_compress(files):
    lvl = choose(["Légère (qualité maximale)", "Recommandée", "Forte (taille minimale)"],
                 "Niveau de compression :", default="Recommandée")
    thr, tgt, q = {"L": (220, 150, 85), "R": (160, 120, 70), "F": (110, 85, 50)}[lvl[0]]
    outs = []
    for f in files:
        doc = open_pdf(f)
        try:
            doc.rewrite_images(dpi_threshold=thr, dpi_target=tgt, quality=q)
        except Exception as e:
            log(f"rewrite_images: {e}")
        try:
            doc.subset_fonts()
        except Exception:
            pass
        dst = out(f, "compressé")
        doc.save(str(dst), garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
        size_note(f, dst)
        outs.append(dst)
    return outs


@action("📄 PDF · Rogner", PDF)
def pdf_crop(files):
    mode = choose(["Visuellement (je dessine la zone à garder)", "Retirer des marges (en mm)",
                   "Supprimer automatiquement les marges blanches"], "Comment rogner ?")
    mm = 72 / 25.4
    outs = []
    for f in files:
        doc = open_pdf(f)
        if mode.startswith("Visuel"):
            sel = pick_regions(doc, "pdf", "Rogner",
                               "Dessine un rectangle autour de la zone à GARDER.\n"
                               "Navigue avec ◀ ▶ pour vérifier les autres pages.")
            pi, [r] = next(iter(sel.items()))
            R = doc[pi].rect
            l, t, rr, b = r[0], r[1], R.width - r[2], R.height - r[3]
            scope = choose(["Toutes les pages", f"Seulement la page {pi + 1}"], "Appliquer le rognage à :")
            targets = range(doc.page_count) if scope.startswith("Toutes") else [pi]
            for p in targets:
                crop_page(doc[p], l, t, rr, b)
        elif mode.startswith("Retirer"):
            v = [float(x) * mm for x in re.split(r"[,; ]+", ask(
                "Marges à retirer en mm : haut, droite, bas, gauche\n(une seule valeur = partout)", "10")) if x]
            t, rr, b, l = (v * 4)[:4] if len(v) == 1 else (v + [0, 0, 0, 0])[:4]
            for pg in doc:
                crop_page(pg, l, t, rr, b)
        else:
            pad = 8
            for pg in doc:
                bb = content_bbox(pg)
                if bb.is_empty:
                    continue
                R = pg.rect
                crop_page(pg, max(0, bb.x0 - pad), max(0, bb.y0 - pad),
                          max(0, R.width - bb.x1 - pad), max(0, R.height - bb.y1 - pad))
        outs.append(save_pdf(doc, out(f, "rogné")))
    return outs


@action("📄 PDF · Convertir en images (JPG / PNG)", PDF)
def pdf_to_images(files):
    fmt = choose(["JPG", "PNG"], "Format des images :")
    dpi = int(choose(["100", "150", "200", "300"], "Résolution (DPI) — 150 écran, 300 impression :", default="150"))
    outs = []
    for f in files:
        doc = open_pdf(f)
        d = outdir(f, "images")
        for i, pg in enumerate(doc, 1):
            pix = pg.get_pixmap(dpi=dpi, alpha=False)
            p = d / f"{Path(f).stem}_page{i:03d}.{fmt.lower()}"
            if fmt == "JPG":
                pix.save(str(p), jpg_quality=90)
            else:
                pix.save(str(p))
        outs.append(d)
    return outs


@action("📄 PDF · Convertir en Word (.docx)", PDF)
def pdf_to_word(files):
    from pdf2docx import Converter
    outs = []
    notify("Conversion en Word en cours…")
    for f in files:
        pw = None
        import pymupdf
        if pymupdf.open(str(f)).needs_pass:
            pw = ask(f"Mot de passe de « {Path(f).name} » :", hidden=True)
        dst = out(f, "word", "docx").with_name(Path(f).stem + ".docx")
        dst = unique(dst)
        cv = Converter(str(f), password=pw)
        try:
            cv.convert(str(dst))
        finally:
            cv.close()
        outs.append(dst)
    return outs


@action("📄 PDF · Convertir en PowerPoint (.pptx)", PDF)
def pdf_to_pptx(files):
    from pptx import Presentation
    from pptx.util import Emu
    outs = []
    for f in files:
        doc = open_pdf(f)
        prs = Presentation()
        R = doc[0].rect
        prs.slide_width, prs.slide_height = Emu(int(R.width * 12700)), Emu(int(R.height * 12700))
        for pg in doc:
            pix = pg.get_pixmap(dpi=170, alpha=False)
            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0, prs.slide_width, prs.slide_height)
        dst = unique(Path(f).with_suffix(".pptx"))
        prs.save(str(dst))
        outs.append(dst)
    NOTES.append("Chaque page devient une diapositive (image).")
    return outs


@action("📄 PDF · Convertir en Excel (tableaux)", PDF)
def pdf_to_excel(files):
    import openpyxl
    outs = []
    for f in files:
        doc = open_pdf(f)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        count = 0
        for pno, pg in enumerate(doc, 1):
            for ti, t in enumerate(pg.find_tables().tables, 1):
                ws = wb.create_sheet(f"Page{pno}_Tableau{ti}"[:31])
                for row in t.extract():
                    ws.append(["" if c is None else str(c).replace("\n", " ") for c in row])
                count += 1
        if not count:
            ws = wb.create_sheet("Texte")
            for pg in doc:
                for line in pg.get_text().splitlines():
                    ws.append(re.split(r"\s{2,}", line.strip()))
            NOTES.append(f"{Path(f).name} : aucun tableau détecté, texte exporté ligne par ligne.")
        else:
            NOTES.append(f"{Path(f).name} : {count} tableau(x) extrait(s).")
        dst = unique(Path(f).with_suffix(".xlsx"))
        wb.save(str(dst))
        outs.append(dst)
    return outs


@action("📄 PDF · Convertir en texte / Markdown", PDF)
def pdf_to_text(files):
    fmt = choose(["Markdown (.md) — garde titres et tableaux", "Texte brut (.txt)"], "Format :")
    outs = []
    for f in files:
        doc = open_pdf(f)
        if fmt.startswith("Markdown"):
            import pymupdf4llm
            txt, e = pymupdf4llm.to_markdown(doc), "md"
        else:
            txt, e = "\n\n".join(p.get_text() for p in doc), "txt"
        dst = unique(Path(f).with_suffix("." + e))
        dst.write_text(txt, encoding="utf-8")
        outs.append(dst)
    return outs


@action("📄 PDF · Extraire les images", PDF)
def pdf_extract_images(files):
    outs = []
    for f in files:
        doc = open_pdf(f)
        d = outdir(f, "images_extraites")
        seen, k = set(), 0
        for pg in doc:
            for img in pg.get_images(full=True):
                x = img[0]
                if x in seen:
                    continue
                seen.add(x)
                info = doc.extract_image(x)
                if not info or info["width"] < 32 or info["height"] < 32:
                    continue
                k += 1
                (d / f"image_{k:03d}.{info['ext']}").write_bytes(info["image"])
        NOTES.append(f"{Path(f).name} : {k} image(s) extraite(s).")
        outs.append(d)
    return outs


@action("🔒 PDF · Protéger par mot de passe", PDF)
def pdf_protect(files):
    import pymupdf
    pw = ask("Mot de passe à définir :", hidden=True)
    if not pw:
        raise ValueError("Mot de passe vide.")
    if ask("Confirme le mot de passe :", hidden=True) != pw:
        raise ValueError("Les mots de passe ne correspondent pas.")
    outs = []
    for f in files:
        doc = open_pdf(f)
        dst = out(f, "protégé")
        doc.save(str(dst), garbage=3, deflate=True, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw=pw, user_pw=pw)
        outs.append(dst)
    NOTES.append("Chiffrement AES-256.")
    return outs


@action("🔓 PDF · Retirer le mot de passe / les restrictions", PDF)
def pdf_unlock(files):
    import pymupdf
    outs = []
    for f in files:
        doc = open_pdf(f)
        if not doc.is_encrypted and not doc.needs_pass:
            NOTES.append(f"{Path(f).name} n'était pas protégé.")
        outs.append(save_pdf(doc, out(f, "déverrouillé"), encryption=pymupdf.PDF_ENCRYPT_NONE))
    return outs


@action("💧 PDF · Ajouter un filigrane", PDF)
def pdf_watermark(files):
    import pymupdf
    text = ask("Texte du filigrane :", "CONFIDENTIEL")
    op = {"Discret": 0.12, "Normal": 0.22, "Marqué": 0.4}[choose(["Discret", "Normal", "Marqué"], "Intensité :", default="Normal")]
    outs = []
    for f in files:
        doc = open_pdf(f)
        for pg in doc:
            R = pg.rect
            diag = math.hypot(R.width, R.height)
            fs = min(diag * 0.7 / max(pymupdf.get_text_length(text, "helv", 1), 1), 160)
            tl = pymupdf.get_text_length(text, "helv", fs)
            ctr = pymupdf.Point(R.width / 2, R.height / 2)
            ang = math.degrees(math.atan2(R.height, R.width))
            pg.insert_text(pymupdf.Point(ctr.x - tl / 2, ctr.y + fs * 0.35), text, fontsize=fs, fontname="helv",
                           color=(0.55, 0.55, 0.55), fill_opacity=op, morph=(ctr, pymupdf.Matrix(ang)), overlay=True)
        outs.append(save_pdf(doc, out(f, "filigrane")))
    return outs


@action("🔢 PDF · Numéroter les pages", PDF)
def pdf_numbers(files):
    import pymupdf
    fmt = choose(["1", "Page 1 / N", "1 / N", "- 1 -"], "Format :", default="Page 1 / N")
    pos = choose(["Bas centre", "Bas droite", "Bas gauche", "Haut droite"], "Position :", default="Bas centre")
    first = int(ask("Commencer la numérotation à partir de la page :", "1"))
    outs = []
    for f in files:
        doc = open_pdf(f)
        N = doc.page_count - first + 1
        for i, pg in enumerate(doc):
            if i + 1 < first:
                continue
            k = i + 2 - first
            s = {"1": f"{k}", "Page 1 / N": f"Page {k} / {N}", "1 / N": f"{k} / {N}", "- 1 -": f"- {k} -"}[fmt]
            R, fs = pg.rect, 10
            tl = pymupdf.get_text_length(s, "helv", fs)
            x = {"centre": (R.width - tl) / 2, "droite": R.width - tl - 36, "gauche": 36}[pos.split()[-1]]
            y = 28 if pos.startswith("Haut") else R.height - 22
            pg.insert_text((x, y), s, fontsize=fs, fontname="helv", color=(0.25, 0.25, 0.25))
        outs.append(save_pdf(doc, out(f, "numéroté")))
    return outs


@action("✍️ PDF · Signer (dessiner ou insérer une signature)", PDF)
def pdf_sign(files):
    opts = (["Utiliser ma signature enregistrée"] if SIG_FILE.exists() else []) + \
        ["Dessiner une signature", "Choisir une image (PNG, JPG…)"]
    m = choose(opts, "Signature :")
    if m.startswith("Utiliser"):
        sig = SIG_FILE
    elif m.startswith("Dessiner"):
        sig = draw_signature()
    else:
        sig = Path(choose_file("Choisis l'image de ta signature", ["public.image"]))
        if ext(sig) in ("heic", "heif", "webp"):
            from PIL import Image
            _heif()
            tmp = Path(tempfile.mkdtemp()) / "sig.png"
            Image.open(sig).save(tmp)
            sig = tmp
    outs = []
    for f in files:
        doc = open_pdf(f)
        sel = pick_regions(doc, "pdf", "Placer la signature",
                           "Dessine un rectangle à l'endroit où placer la signature.\n"
                           "Tu peux en placer plusieurs (sur différentes pages : ◀ ▶).",
                           multi=True, start_page=doc.page_count - 1)
        for pi, rects in sel.items():
            for r in rects:
                doc[pi].insert_image(r, filename=str(sig), keep_proportion=True, overlay=True)
        outs.append(save_pdf(doc, out(f, "signé")))
    return outs


@action("⬛ PDF · Caviarder (masquer définitivement)", PDF)
def pdf_redact(files):
    m = choose(["Dessiner les zones à masquer", "Rechercher un mot / une expression"], "Caviarder :")
    words = ask("Mots ou expressions à masquer (séparés par ; ) :", "") if m.startswith("Rechercher") else None
    outs = []
    for f in files:
        doc = open_pdf(f)
        hits = 0
        if words is None:
            sel = pick_regions(doc, "pdf", "Caviarder",
                               "Dessine un rectangle sur chaque zone à masquer (plusieurs possibles, sur toutes les pages).",
                               multi=True)
            for pi, rects in sel.items():
                for r in rects:
                    doc[pi].add_redact_annot(r, fill=(0, 0, 0))
                    hits += 1
        else:
            for pg in doc:
                for w in [w.strip() for w in words.split(";") if w.strip()]:
                    for r in pg.search_for(w):
                        pg.add_redact_annot(r, fill=(0, 0, 0))
                        hits += 1
        for pg in doc:
            pg.apply_redactions()
        NOTES.append(f"{Path(f).name} : {hits} zone(s) masquée(s) — le contenu dessous est réellement supprimé.")
        outs.append(save_pdf(doc, out(f, "caviardé"), garbage=4))
    return outs


def tess_langs():
    pref = os.environ.get("TESSDATA_PREFIX")
    if pref and Path(pref).is_dir():
        found = [x.stem for x in Path(pref).glob("*.traineddata")]
        if found:
            return found
    t = which("tesseract", "/opt/homebrew/bin/tesseract", "/usr/local/bin/tesseract")
    if not t:
        return None
    try:
        r = subprocess.run([t, "--list-langs"], capture_output=True, text=True)
        return [l.strip() for l in (r.stdout + r.stderr).splitlines()[1:] if l.strip() and " " not in l.strip()]
    except Exception:
        return None


@action("🔍 PDF · OCR — rendre un scan cherchable", PDF)
def pdf_ocr(files):
    import pymupdf
    langs = tess_langs()
    if langs is None:
        raise RuntimeError("L'OCR nécessite Tesseract. Installe-le avec : brew install tesseract")
    lang = "+".join(l for l in ("fra", "eng") if l in langs) or "eng"
    outs = []
    notify("OCR en cours… (quelques secondes par page)")
    for f in files:
        doc = open_pdf(f)
        od = pymupdf.open()
        for pg in doc:
            pix = pg.get_pixmap(dpi=300)
            od.insert_pdf(pymupdf.open("pdf", pix.pdfocr_tobytes(language=lang)))
        outs.append(save_pdf(od, out(f, "ocr"), garbage=4))
    NOTES.append(f"Langue(s) OCR : {lang}. Le texte est maintenant sélectionnable et cherchable.")
    return outs


@action("📄 PDF · Aplatir (formulaires et annotations)", PDF)
def pdf_flatten(files):
    outs = []
    for f in files:
        doc = open_pdf(f)
        doc.bake(annots=True, widgets=True)
        outs.append(save_pdf(doc, out(f, "aplati")))
    return outs


@action("📄 PDF · Noir et blanc (niveaux de gris)", PDF)
def pdf_gray(files):
    import pymupdf
    outs = []
    for f in files:
        doc = open_pdf(f)
        od = pymupdf.open()
        for pg in doc:
            pix = pg.get_pixmap(dpi=150, colorspace=pymupdf.csGRAY)
            np_ = od.new_page(width=pg.rect.width, height=pg.rect.height)
            np_.insert_image(np_.rect, pixmap=pix)
        outs.append(save_pdf(od, out(f, "gris"), garbage=4))
    NOTES.append("Les pages sont converties en images (texte non sélectionnable — utilise l'OCR si besoin).")
    return outs


@action("🏷️ PDF · Métadonnées (titre, auteur, nettoyage)", PDF)
def pdf_meta(files):
    m = choose(["Modifier titre / auteur / sujet", "Supprimer toutes les métadonnées"], "Métadonnées :")
    outs = []
    for f in files:
        doc = open_pdf(f)
        if m.startswith("Supprimer"):
            doc.set_metadata({})
            doc.del_xml_metadata()
        else:
            md = doc.metadata or {}
            md["title"] = ask(f"Titre de « {Path(f).name} » :", md.get("title") or Path(f).stem)
            md["author"] = ask("Auteur :", md.get("author") or "")
            md["subject"] = ask("Sujet :", md.get("subject") or "")
            doc.set_metadata(md)
        outs.append(save_pdf(doc, out(f, "meta")))
    return outs


@action("🛠️ PDF · Réparer un PDF endommagé", PDF)
def pdf_repair(files):
    outs = []
    for f in files:
        doc = open_pdf(f)
        outs.append(save_pdf(doc, out(f, "réparé"), garbage=4, clean=True))
    return outs


@action("ℹ️ PDF · Informations", PDF)
def pdf_info(files):
    for f in files:
        doc = open_pdf(f)
        R = doc[0].rect
        md = doc.metadata or {}
        nimg = sum(len(p.get_images()) for p in doc)
        txt = sum(len(p.get_text().strip()) > 20 for p in doc)
        lines = [
            f"Pages : {doc.page_count}", f"Taille : {human(Path(f).stat().st_size)}",
            f"Format page 1 : {R.width / 72 * 25.4:.0f} × {R.height / 72 * 25.4:.0f} mm",
            f"Images : {nimg}", f"Pages avec texte : {txt}/{doc.page_count}" + ("  → probablement un scan (OCR conseillé)" if txt == 0 else ""),
            f"Chiffré : {'oui' if doc.is_encrypted else 'non'}", f"Formulaire : {'oui' if doc.is_form_pdf else 'non'}",
            f"Titre : {md.get('title') or '—'}", f"Auteur : {md.get('author') or '—'}",
            f"Créé avec : {md.get('creator') or md.get('producer') or '—'}",
        ]
        buttons(f"{Path(f).name}\n\n" + "\n".join(lines), ["OK"])
    return []

# ───────────────────────── IMAGES ─────────────────────────

def load_img(f):
    from PIL import Image, ImageOps
    _heif()
    im = Image.open(f)
    im.load()
    im = ImageOps.exif_transpose(im) or im
    return im


def flatten(im, bg=(255, 255, 255)):
    from PIL import Image
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        b = Image.new("RGB", im.size, bg)
        b.paste(im, mask=im.getchannel("A"))
        return b
    return im if im.mode in ("RGB", "L") else im.convert("RGB")


PIL_FMT = {"jpg": "JPEG", "jpeg": "JPEG", "png": "PNG", "webp": "WEBP", "heic": "HEIF", "heif": "HEIF",
           "avif": "AVIF", "tif": "TIFF", "tiff": "TIFF", "gif": "GIF", "bmp": "BMP", "ico": "ICO"}


def save_img(im, dst, quality=90, keep_exif=True):
    from PIL import Image
    e = ext(dst)
    fmt = PIL_FMT.get(e, "PNG")
    kw = {}
    exif = im.info.get("exif") if keep_exif else None
    if im.mode not in ("RGB", "RGBA", "L", "LA", "P", "1"):
        im = im.convert("RGBA" if "A" in im.getbands() else "RGB")
    if fmt in ("JPEG", "BMP"):
        im = flatten(im)
        if fmt == "JPEG":
            kw = dict(quality=quality, optimize=True, progressive=True)
    elif fmt == "PNG":
        kw = dict(optimize=True)
    elif fmt in ("WEBP", "AVIF", "HEIF"):
        if im.mode not in ("RGB", "RGBA"):
            im = im.convert("RGBA" if im.mode in ("LA", "P") else "RGB")
        kw = dict(quality=quality)
        if fmt == "WEBP":
            kw["method"] = 6
    elif fmt == "TIFF":
        kw = dict(compression="tiff_lzw")
    elif fmt == "ICO":
        kw = dict(sizes=[(s, s) for s in (256, 128, 64, 48, 32, 16)])
    if exif and fmt in ("JPEG", "WEBP", "PNG", "HEIF", "AVIF", "TIFF"):
        kw["exif"] = exif
    if not keep_exif:
        icc = im.info.get("icc_profile")
        im.info = {"icc_profile": icc} if icc else {}
    im.save(str(dst), fmt, **kw)
    return Path(dst)


def writable_formats():
    from PIL import Image
    _heif()
    Image.init()
    fm = ["JPG", "PNG", "WEBP", "HEIC", "AVIF", "TIFF", "GIF", "BMP", "ICO"]
    return [x for x in fm if PIL_FMT[x.lower()] in Image.SAVE]


@action("🖼️ Image · Convertir le format", IMG)
def img_convert(files):
    fmt = choose(writable_formats(), "Convertir en :").lower()
    outs = []
    for f in files:
        im = load_img(f)
        if ext(f) == "gif" and fmt == "gif":
            continue
        outs.append(save_img(im, unique(Path(f).with_suffix("." + fmt)), quality=92))
    return outs


@action("🖼️ Image · Redimensionner", IMG)
def img_resize(files):
    from PIL import Image
    s = ask("Nouvelle taille :\n• 50%  → pourcentage\n• 1920 → plus grand côté en pixels\n• 1080x1080 → tenir dans un cadre", "50%")
    outs = []
    for f in files:
        im = load_img(f)
        w, h = im.size
        if s.endswith("%"):
            k = float(s[:-1]) / 100
            nw, nh = round(w * k), round(h * k)
        elif "x" in s.lower():
            bw, bh = [int(x) for x in s.lower().split("x")]
            k = min(bw / w, bh / h)
            nw, nh = round(w * k), round(h * k)
        else:
            k = int(s) / max(w, h)
            nw, nh = round(w * k), round(h * k)
        r = im.resize((max(1, nw), max(1, nh)), Image.LANCZOS)
        r.info = im.info
        outs.append(save_img(r, out(f, f"{nw}x{nh}"), quality=92))
    return outs


@action("🖼️ Image · Compresser", IMG)
def img_compress(files):
    from PIL import Image
    lvl = choose(["Légère", "Recommandée", "Forte"], "Niveau de compression :", default="Recommandée")
    q, maxs = {"L": (85, None), "R": (72, 3000), "F": (55, 2000)}[lvl[0]]
    outs = []
    for f in files:
        im = load_img(f)
        if maxs and max(im.size) > maxs:
            info = im.info
            im.thumbnail((maxs, maxs), Image.LANCZOS)
            im.info = info
        e = ext(f)
        if e in ("png", "gif", "bmp", "tif", "tiff") and lvl[0] == "F" and e != "gif":
            if "A" in im.getbands():
                im = im.quantize(256, method=Image.Quantize.FASTOCTREE)
            else:
                im = im.convert("RGB").quantize(256)
        dst = out(f, "compressé", "jpg" if e in ("bmp", "tif", "tiff") else None)
        save_img(im, dst, quality=q)
        size_note(f, dst)
        outs.append(dst)
    return outs


@action("✂️ Image · Rogner", IMG)
def img_crop(files):
    from PIL import Image, ImageChops
    mode = choose(["Visuellement (je dessine la zone)", "Format imposé (1:1, 16:9, 4:5…)",
                   "Supprimer automatiquement les bords unis"], "Comment rogner ?")
    ratio = None
    if mode.startswith("Format"):
        r = choose(["1:1 (carré)", "4:5 (Instagram portrait)", "16:9 (paysage)", "9:16 (story)",
                    "4:3", "3:2", "3:4", "2:3"], "Format :")
        a, b = r.split()[0].split(":")
        ratio = int(a) / int(b)
    outs = []
    for f in files:
        im = load_img(f)
        w, h = im.size
        if mode.startswith("Visuel"):
            sel = pick_regions(im, "img", f"Rogner — {Path(f).name}", "Dessine un rectangle autour de la zone à GARDER.")
            box = tuple(round(v) for v in sel[0][0])
        elif ratio:
            if w / h > ratio:
                nw = round(h * ratio)
                box = ((w - nw) // 2, 0, (w - nw) // 2 + nw, h)
            else:
                nh = round(w / ratio)
                box = (0, (h - nh) // 2, w, (h - nh) // 2 + nh)
        else:
            box = None
            if "A" in im.getbands() and im.getchannel("A").getpixel((0, 0)) < 10:
                box = im.getchannel("A").getbbox()
            if not box:
                rgb = im.convert("RGB")
                diff = ImageChops.difference(rgb, Image.new("RGB", im.size, rgb.getpixel((0, 0))))
                box = ImageChops.add(diff, diff, 2.0, -24).getbbox()
            if not box:
                NOTES.append(f"{Path(f).name} : aucun bord détecté.")
                continue
        c = im.crop(box)
        c.info = im.info
        outs.append(save_img(c, out(f, "rogné"), quality=93))
    return outs


@action("🔄 Image · Pivoter / retourner", IMG)
def img_rotate(files):
    from PIL import Image
    T = Image.Transpose
    m = choose(["90° vers la droite", "90° vers la gauche", "180°", "Miroir horizontal", "Miroir vertical"], "Transformation :")
    op = {"90° vers la droite": T.ROTATE_270, "90° vers la gauche": T.ROTATE_90, "180°": T.ROTATE_180,
          "Miroir horizontal": T.FLIP_LEFT_RIGHT, "Miroir vertical": T.FLIP_TOP_BOTTOM}[m]
    outs = []
    for f in files:
        im = load_img(f)
        r = im.transpose(op)
        r.info = im.info
        outs.append(save_img(r, out(f, "pivoté"), quality=93))
    return outs


@action("🕵️ Image · Supprimer les métadonnées (EXIF, GPS)", IMG)
def img_strip(files):
    outs = []
    for f in files:
        im = load_img(f)
        outs.append(save_img(im, out(f, "sans-meta"), quality=95, keep_exif=False))
    NOTES.append("Position GPS, appareil, date et autres données EXIF retirées.")
    return outs


@action("📄 Image · Convertir en PDF", IMG)
def img_to_pdf(files):
    import pymupdf
    one = len(files) == 1 or choose(["Un seul PDF avec toutes les images", "Un PDF par image"], "Regrouper ?").startswith("Un seul")
    fit = choose(["A4 (ajusté avec marges)", "Taille de l'image"], "Format des pages :").startswith("A4")
    if one:
        d = pymupdf.open()
        for f in files:
            add_image_page(d, f, fit)
        name = "Images.pdf" if len(files) > 1 else Path(files[0]).stem + ".pdf"
        return [save_pdf(d, unique(Path(files[0]).with_name(name)), garbage=4)]
    outs = []
    for f in files:
        d = pymupdf.open()
        add_image_page(d, f, fit)
        outs.append(save_pdf(d, unique(Path(f).with_suffix(".pdf")), garbage=4))
    return outs


@action("💧 Image · Ajouter un filigrane / copyright", IMG)
def img_watermark(files):
    from PIL import Image, ImageDraw, ImageFont
    text = ask("Texte du filigrane :", f"© {datetime.date.today().year}")
    pos = choose(["Bas droite", "Bas gauche", "Centre (grand)", "Mosaïque (répété en diagonale)"], "Position :")
    outs = []
    for f in files:
        im = load_img(f)
        base = im.convert("RGBA")
        W, H = base.size
        fs = max(14, int(min(W, H) * (0.12 if pos.startswith("Centre") else 0.045)))
        try:
            font = ImageFont.truetype(font_file(), fs)
        except Exception:
            font = ImageFont.load_default(size=fs)
        ov = Image.new("RGBA", base.size, (0, 0, 0, 0))
        d = ImageDraw.Draw(ov)
        x0, y0, x1, y1 = d.textbbox((0, 0), text, font=font)
        tw, th = x1 - x0, y1 - y0
        if pos.startswith("Mosa"):
            tile = Image.new("RGBA", (tw + fs * 3, th + fs * 3), (0, 0, 0, 0))
            ImageDraw.Draw(tile).text((fs, fs), text, font=font, fill=(255, 255, 255, 80))
            tile = tile.rotate(30, expand=True)
            for yy in range(-tile.height, H + tile.height, tile.height):
                for xx in range(-tile.width, W + tile.width, tile.width):
                    ov.paste(tile, (xx, yy), tile)
        else:
            m = fs
            p = {"Bas droite": (W - tw - m, H - th - m * 1.4), "Bas gauche": (m, H - th - m * 1.4),
                 "Centre (grand)": ((W - tw) / 2, (H - th) / 2)}[pos]
            d.text((p[0] + 2, p[1] + 2), text, font=font, fill=(0, 0, 0, 90))
            d.text(p, text, font=font, fill=(255, 255, 255, 190))
        res = Image.alpha_composite(base, ov)
        res.info = im.info
        if "A" not in im.getbands():
            res = res.convert("RGB")
            res.info = im.info
        outs.append(save_img(res, out(f, "filigrane"), quality=93))
    return outs


@action("🎞️ Images · Créer un GIF animé", IMG, minf=2)
def img_gif(files):
    from PIL import Image
    ms = int(ask("Durée de chaque image (millisecondes) :", "600"))
    width = int(ask("Largeur du GIF (pixels) :", "800"))
    frames = []
    for f in files:
        im = load_img(f).convert("RGB")
        im.thumbnail((width, width * 4), Image.LANCZOS)
        frames.append(im)
    W, H = max(i.width for i in frames), max(i.height for i in frames)
    canv = []
    for im in frames:
        c = Image.new("RGB", (W, H), "white")
        c.paste(im, ((W - im.width) // 2, (H - im.height) // 2))
        canv.append(c)
    dst = unique(Path(files[0]).with_name("Animation.gif"))
    canv[0].save(dst, save_all=True, append_images=canv[1:], duration=ms, loop=0, optimize=True)
    return [dst]

# ───────────────────────── VIDÉO ─────────────────────────

def _vlabel(files):
    notify(f"Traitement vidéo en cours ({len(files)} fichier(s))… je te préviens à la fin.")


@action("🎬 Vidéo · Convertir en MP4 (compatible partout)", VID)
def vid_mp4(files):
    _vlabel(files)
    outs = []
    for f in files:
        dst = unique(Path(f).with_name(Path(f).stem + (".mp4" if ext(f) != "mp4" else "_h264.mp4")))
        ff("-i", f, "-c:v", "libx264", "-crf", "21", "-preset", "medium", "-pix_fmt", "yuv420p",
           "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart", dst)
        outs.append(dst)
    return outs


@action("🎬 Vidéo · Compresser", VID)
def vid_compress(files):
    lvl = choose(["Légère (qualité quasi identique)", "Recommandée (max 1080p)", "Forte (max 720p, idéal mail)"],
                 "Niveau :", default="Recommandée (max 1080p)")
    crf, hmax = {"L": (24, None), "R": (27, 1080), "F": (31, 720)}[lvl[0]]
    _vlabel(files)
    outs = []
    for f in files:
        dst = out(f, "compressé", "mp4")
        vf = ["-vf", f"scale=-2:'min({hmax},ih)'"] if hmax else []
        ff("-i", f, *vf, "-c:v", "libx264", "-crf", crf, "-preset", "medium", "-pix_fmt", "yuv420p",
           "-c:a", "aac", "-b:a", "128k", "-movflags", "+faststart", dst)
        size_note(f, dst)
        outs.append(dst)
    return outs


@action("🎬 Vidéo · Couper (garder un extrait)", VID)
def vid_trim(files):
    a = parse_time(ask("Début de l'extrait (ex : 1:23 ou 83 ou 0:01:23) :", "0:00"))
    b = parse_time(ask("Fin de l'extrait (vide = jusqu'à la fin) :", ""))
    _vlabel(files)
    outs = []
    for f in files:
        dst = out(f, "extrait", "mp4")
        to = ["-to", b] if b else []
        ff("-ss", a or 0, *to, "-i", f, "-c:v", "libx264", "-crf", "20", "-preset", "fast", "-pix_fmt", "yuv420p",
           "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart", dst)
        outs.append(dst)
    return outs


@action("🎬 Vidéo · Convertir en GIF", VID)
def vid_gif(files):
    w = int(choose(["320", "480", "640", "800"], "Largeur du GIF (px) :", default="480"))
    fps = int(choose(["10", "12", "15", "20"], "Images par seconde :", default="12"))
    a = parse_time(ask("Début (ex : 0:05, vide = début) :", ""))
    dur = parse_time(ask("Durée en secondes (vide = tout, attention aux gros fichiers) :", "6"))
    _vlabel(files)
    outs = []
    for f in files:
        dst = out(f, "anim", "gif")
        pre = (["-ss", a] if a else []) + (["-t", dur] if dur else [])
        ff(*pre, "-i", f, "-vf",
           f"fps={fps},scale={w}:-1:flags=lanczos,split[s0][s1];[s0]palettegen=stats_mode=diff[p];[s1][p]paletteuse=dither=bayer:bayer_scale=4",
           "-loop", "0", dst)
        outs.append(dst)
    return outs


@action("🎬 Vidéo · Extraire l'audio (MP3)", VID)
def vid_audio(files):
    _vlabel(files)
    outs = []
    for f in files:
        dst = unique(Path(f).with_suffix(".mp3"))
        ff("-i", f, "-vn", "-c:a", "libmp3lame", "-q:a", "2", dst)
        outs.append(dst)
    return outs


@action("🎬 Vidéo · Retirer le son", VID)
def vid_mute(files):
    outs = []
    for f in files:
        dst = out(f, "muet")
        ff("-i", f, "-c:v", "copy", "-an", dst)
        outs.append(dst)
    return outs


@action("🎬 Vidéo · Redimensionner (1080p / 720p / 480p / vertical)", VID)
def vid_resize(files):
    m = choose(["1080p", "720p", "480p", "Vertical 1080×1920 (recadré au centre)", "Carré 1080×1080 (recadré)"], "Taille :")
    vf = {"1080p": "scale=-2:1080", "720p": "scale=-2:720", "480p": "scale=-2:480",
          "Vertical 1080×1920 (recadré au centre)": "scale=1080:1920:force_original_aspect_ratio=increase,crop=1080:1920",
          "Carré 1080×1080 (recadré)": "scale=1080:1080:force_original_aspect_ratio=increase,crop=1080:1080"}[m]
    _vlabel(files)
    outs = []
    for f in files:
        dst = out(f, m.split()[0].lower(), "mp4")
        ff("-i", f, "-vf", vf, "-c:v", "libx264", "-crf", "21", "-preset", "medium", "-pix_fmt", "yuv420p",
           "-c:a", "aac", "-b:a", "160k", "-movflags", "+faststart", dst)
        outs.append(dst)
    return outs


@action("🎬 Vidéo · Pivoter", VID)
def vid_rotate(files):
    m = choose(["90° vers la droite", "90° vers la gauche", "180°"], "Rotation :")
    vf = {"90° vers la droite": "transpose=1", "90° vers la gauche": "transpose=2", "180°": "transpose=1,transpose=1"}[m]
    _vlabel(files)
    outs = []
    for f in files:
        dst = out(f, "pivoté", "mp4")
        ff("-i", f, "-vf", vf, "-c:v", "libx264", "-crf", "20", "-preset", "medium", "-pix_fmt", "yuv420p",
           "-c:a", "copy", "-movflags", "+faststart", dst)
        outs.append(dst)
    return outs


@action("🎬 Vidéo · Capturer une image", VID)
def vid_frame(files):
    t = parse_time(ask("À quel moment ? (ex : 0:12) :", "0:01")) or 0
    outs = []
    for f in files:
        dst = out(f, f"image_{int(t)}s", "jpg")
        ff("-ss", t, "-i", f, "-frames:v", "1", "-q:v", "2", dst)
        outs.append(dst)
    return outs

# ───────────────────────── AUDIO ─────────────────────────

ACODEC = {"mp3": ["-c:a", "libmp3lame", "-q:a", "2"], "m4a": ["-c:a", "aac", "-b:a", "192k"],
          "wav": ["-c:a", "pcm_s16le"], "flac": ["-c:a", "flac"], "ogg": ["-c:a", "libvorbis", "-q:a", "5"],
          "aiff": ["-c:a", "pcm_s16be"]}


@action("🎵 Audio · Convertir le format", AUD | VID)
def aud_convert(files):
    e = choose(["MP3", "M4A (AAC)", "WAV", "FLAC", "OGG", "AIFF"], "Convertir en :").split()[0].lower()
    outs = []
    for f in files:
        dst = unique(Path(f).with_suffix("." + e))
        ff("-i", f, "-vn", *ACODEC[e], dst)
        outs.append(dst)
    return outs


@action("🎵 Audio · Compresser", AUD)
def aud_compress(files):
    br = choose(["192k (très bonne qualité)", "128k (standard)", "96k (voix / podcast)", "64k mono (voix, très léger)"],
                "Débit :", default="128k (standard)")
    outs = []
    for f in files:
        dst = out(f, "compressé", "mp3")
        mono = ["-ac", "1"] if "mono" in br else []
        ff("-i", f, "-vn", *mono, "-c:a", "libmp3lame", "-b:a", br.split()[0], dst)
        size_note(f, dst)
        outs.append(dst)
    return outs


@action("🎵 Audio · Couper (garder un extrait)", AUD)
def aud_trim(files):
    a = parse_time(ask("Début (ex : 1:23) :", "0:00")) or 0
    b = parse_time(ask("Fin (vide = jusqu'à la fin) :", ""))
    fade = choose(["Sans fondu", "Fondu entrée + sortie (2 s)"], "Fondu :").startswith("Fondu")
    outs = []
    for f in files:
        dst = out(f, "extrait")
        to = ["-to", b] if b else []
        af = []
        if fade:
            dur = (b - a) if b else None
            fl = "afade=t=in:d=2"
            if dur:
                fl += f",afade=t=out:st={max(0, dur - 2)}:d=2"
            af = ["-af", fl]
        ff("-ss", a, *to, "-i", f, "-vn", *af, *ACODEC.get(ext(f), ["-c:a", "libmp3lame", "-q:a", "2"]), dst)
        outs.append(dst)
    return outs


@action("🎵 Audio · Normaliser le volume", AUD | VID)
def aud_norm(files):
    outs = []
    for f in files:
        if ext(f) in VID:
            dst = out(f, "volume-normalisé")
            ff("-i", f, "-c:v", "copy", "-af", "loudnorm=I=-16:TP=-1.5:LRA=11", "-c:a", "aac", "-b:a", "192k", dst)
        else:
            dst = out(f, "volume-normalisé")
            ff("-i", f, "-vn", "-af", "loudnorm=I=-16:TP=-1.5:LRA=11", *ACODEC.get(ext(f), ACODEC["mp3"]), dst)
        outs.append(dst)
    return outs


@action("🎵 Audio · Fusionner (mettre bout à bout)", AUD, minf=2)
def aud_merge(files):
    dst = unique(Path(files[0]).with_name("Fusion.mp3"))
    args = []
    for f in files:
        args += ["-i", f]
    n = len(files)
    ff(*args, "-filter_complex", "".join(f"[{i}:a]" for i in range(n)) + f"concat=n={n}:v=0:a=1[a]",
       "-map", "[a]", "-c:a", "libmp3lame", "-q:a", "2", dst)
    return [dst]

# ───────────────────────── DOCUMENTS ─────────────────────────

CSS = ("body{font-family:sans-serif;font-size:11pt;line-height:1.45;color:#111}"
       "h1{font-size:21pt}h2{font-size:16pt}h3{font-size:13pt}code,pre{font-family:monospace;font-size:9pt}"
       "table{border-collapse:collapse}td,th{border:1px solid #888;padding:3px 6px}th{background:#eee}"
       "blockquote{color:#555;margin-left:14pt}")


def to_html(f):
    import html as H
    e = ext(f)
    raw = Path(f).read_text(encoding="utf-8", errors="replace")
    if e in ("md", "markdown"):
        import markdown
        return markdown.markdown(raw, extensions=["tables", "fenced_code", "sane_lists"])
    if e in ("html", "htm"):
        return raw
    return "<div style='font-family:monospace;font-size:10pt'>" + H.escape(raw).replace("\n", "<br/>") + "</div>"


def html_to_pdf(html, dst, base=None):
    import pymupdf
    arch = pymupdf.Archive(str(base)) if base else None
    story = pymupdf.Story(html=html, user_css=CSS, archive=arch)
    w = pymupdf.DocumentWriter(str(dst))
    mb = pymupdf.paper_rect("a4")
    where = mb + (56, 56, -56, -56)
    more = 1
    while more:
        dev = w.begin_page(mb)
        more, _ = story.place(where)
        story.draw(dev)
        w.end_page()
    w.close()
    return dst


def soffice():
    return which("soffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                 r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe")


def word_app():
    return os.path.exists("/Applications/Microsoft Word.app")


def lo_convert(src, fmt, dst):
    """Conversion via LibreOffice (headless). Retourne True si réussi."""
    so = soffice()
    if not so:
        return False
    tmp = Path(tempfile.mkdtemp())
    prof = Path(tempfile.mkdtemp())
    try:
        subprocess.run([so, f"-env:UserInstallation=file://{prof}", "--headless", "--norestore",
                        "--convert-to", fmt, "--outdir", str(tmp), str(src)],
                       capture_output=True, timeout=600)
        res = next(tmp.glob("*." + fmt.split(":")[0]), None)
        if not res or res.stat().st_size == 0:
            return False
        shutil.move(str(res), dst)
        return True
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
        shutil.rmtree(prof, ignore_errors=True)


def word_com_to_pdf(src, dst):
    """Windows : conversion fidèle via Microsoft Word s'il est installé."""
    try:
        import pythoncom
        import win32com.client
    except Exception:
        return False
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(Path(src).resolve()), ReadOnly=True, AddToRecentFiles=False)
        doc.SaveAs2(str(Path(dst).resolve()), FileFormat=17)
        doc.Close(False)
        return Path(dst).exists()
    except Exception as err:
        log(f"Word COM : {err}")
        return False
    finally:
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def textutil_to_pdf(src, dst):
    """Secours sans LibreOffice : moteur de conversion intégré à macOS (textutil) puis mise en page A4."""
    tmp = Path(tempfile.mkdtemp())
    try:
        html = tmp / "doc.html"
        r = subprocess.run(["/usr/bin/textutil", "-convert", "html", "-output", str(html), str(src)],
                           capture_output=True, text=True)
        if r.returncode != 0 or not html.exists():
            raise RuntimeError("textutil : " + r.stderr.strip())
        h = html.read_text(encoding="utf-8", errors="replace")
        h = re.sub(r"<style[^>]*>.*?</style>", "", h, flags=re.S | re.I)
        html_to_pdf(h, dst, tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@action("📝 Document · Convertir en PDF", OFFICE | TEXT)
def doc_to_pdf(files):
    outs = []
    for f in files:
        e = ext(f)
        dst = unique(Path(f).with_suffix(".pdf"))
        if e in TEXT:
            html_to_pdf(to_html(f), dst, Path(f).parent)
        elif soffice():
            notify("Conversion en PDF en cours…")
            if not lo_convert(f, "pdf", dst):
                raise RuntimeError(f"LibreOffice n'a pas pu convertir {Path(f).name}.")
        elif e in WORDS and IS_WIN and word_com_to_pdf(f, dst):
            pass
        elif e in WORDS:
            if IS_MAC:
                textutil_to_pdf(f, dst)
            elif e == "docx":
                import mammoth
                with open(f, "rb") as fh:
                    html_to_pdf(mammoth.convert_to_html(fh).value, dst)
            else:
                raise RuntimeError(f"Pour convertir les fichiers .{e} en PDF, installe LibreOffice ou Microsoft Word.")
            NOTES.append("Conversion simplifiée (installe LibreOffice pour une mise en page fidèle).")
        else:
            raise RuntimeError(f"Pour convertir les fichiers .{e} en PDF, installe LibreOffice (gratuit) :\n"
                               "brew install --cask libreoffice\nPuis relance l'action.")
        outs.append(dst)
    return outs


@action("📝 Office · Convertir (Word ⇄ ODT, PowerPoint, Excel…)", OFFICE)
def office_convert(files):
    if not soffice():
        raise RuntimeError("Cette action nécessite LibreOffice : brew install --cask libreoffice")
    groups = {"docx": ["docx", "odt", "rtf", "txt", "html"], "pptx": ["pptx", "odp"], "xlsx": ["xlsx", "ods", "csv"]}
    fam = "pptx" if ext(files[0]) in ("pptx", "ppt", "odp") else "xlsx" if ext(files[0]) in ("xlsx", "xls", "ods") else "docx"
    tgt = choose([x.upper() for x in groups[fam] if x != ext(files[0])], "Convertir en :").lower()
    outs = []
    notify("Conversion en cours…")
    for f in files:
        dst = unique(Path(f).with_suffix("." + tgt))
        fmt = "txt:Text (encoded):UTF8" if tgt == "txt" else tgt
        if not lo_convert(f, fmt, dst):
            raise RuntimeError(f"Conversion impossible pour {Path(f).name}.")
        outs.append(dst)
    return outs


@action("📝 Word · Convertir en Markdown / HTML / Texte", {"docx"})
def docx_convert(files):
    import mammoth
    fmt = choose(["Markdown (.md)", "HTML (.html)", "Texte brut (.txt)"], "Convertir en :")
    outs = []
    for f in files:
        with open(f, "rb") as fh:
            if fmt.startswith("Markdown"):
                txt, e = mammoth.convert_to_markdown(fh).value, "md"
            elif fmt.startswith("HTML"):
                body = mammoth.convert_to_html(fh).value
                txt, e = f"<!doctype html><meta charset='utf-8'><style>{CSS}</style><body>{body}</body>", "html"
            else:
                txt, e = mammoth.extract_raw_text(fh).value, "txt"
        dst = unique(Path(f).with_suffix("." + e))
        dst.write_text(txt, encoding="utf-8")
        outs.append(dst)
    return outs


@action("📝 Markdown · Convertir en HTML", {"md", "markdown"})
def md_html(files):
    outs = []
    for f in files:
        dst = unique(Path(f).with_suffix(".html"))
        dst.write_text(f"<!doctype html><meta charset='utf-8'><title>{Path(f).stem}</title>"
                       f"<style>body{{max-width:46em;margin:3em auto;padding:0 1em}}{CSS}</style>{to_html(f)}", encoding="utf-8")
        outs.append(dst)
    return outs

# ───────────────────────── TABLEURS ─────────────────────────

def read_tables(f):
    import pandas as pd
    e = ext(f)
    if e in ("csv", "tsv"):
        return {Path(f).stem: pd.read_csv(f, sep="\t" if e == "tsv" else None, engine="python", encoding_errors="replace")}
    if e == "json":
        data = json.loads(Path(f).read_text(encoding="utf-8"))
        if isinstance(data, dict) and all(isinstance(v, list) for v in data.values()):
            return {k: pd.json_normalize(v) for k, v in data.items()}
        return {Path(f).stem: pd.json_normalize(data if isinstance(data, list) else [data])}
    return pd.read_excel(f, sheet_name=None)


@action("📊 Tableur · Convertir (Excel ⇄ CSV ⇄ JSON)", TAB)
def tab_convert(files):
    import pandas as pd
    tgt = choose(["Excel (.xlsx)", "CSV", "JSON"], "Convertir en :")
    sep = ";"
    if tgt == "CSV":
        sep = ";" if choose(["Point-virgule ; (Excel en français)", "Virgule ,"], "Séparateur :").startswith("Point") else ","
    outs = []
    for f in files:
        sheets = read_tables(f)
        if tgt.startswith("Excel"):
            dst = unique(Path(f).with_suffix(".xlsx"))
            with pd.ExcelWriter(dst) as xw:
                for name, df in sheets.items():
                    df.to_excel(xw, sheet_name=str(name)[:31] or "Feuille1", index=False)
            outs.append(dst)
        else:
            for name, df in sheets.items():
                suffix = f"_{name}" if len(sheets) > 1 else ""
                dst = unique(Path(f).with_name(f"{Path(f).stem}{suffix}." + ("csv" if tgt == "CSV" else "json")))
                if tgt == "CSV":
                    df.to_csv(dst, sep=sep, index=False, encoding="utf-8-sig")
                else:
                    dst.write_text(df.to_json(orient="records", force_ascii=False, indent=2, date_format="iso"), encoding="utf-8")
                outs.append(dst)
    return outs


@action("📊 Tableur · Fusionner plusieurs fichiers en un Excel", TAB, minf=2)
def tab_merge(files):
    import pandas as pd
    mode = choose(["Une feuille par fichier", "Tout empiler dans une seule feuille"], "Fusion :")
    dst = unique(Path(files[0]).with_name("Fusion.xlsx"))
    frames = []
    with pd.ExcelWriter(dst) as xw:
        for f in files:
            for name, df in read_tables(f).items():
                if mode.startswith("Une"):
                    df.to_excel(xw, sheet_name=f"{Path(f).stem}"[:31], index=False)
                else:
                    df.insert(0, "Source", Path(f).name)
                    frames.append(df)
        if frames:
            pd.concat(frames, ignore_index=True).to_excel(xw, sheet_name="Fusion", index=False)
    return [dst]

# ───────────────────────── GÉNÉRAL ─────────────────────────

@action("📦 Compresser en ZIP")
def make_zip(files):
    p0 = Path(files[0])
    dst = unique(p0.parent / ("Archive.zip" if len(files) > 1 else p0.name + ".zip"))
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for f in files:
            p = Path(f)
            if p.is_dir():
                for s in p.rglob("*"):
                    if s.is_file() and s.name != ".DS_Store":
                        z.write(s, s.relative_to(p.parent))
            else:
                z.write(p, p.name)
    tot = sum((s.stat().st_size for f in files for s in ([Path(f)] if Path(f).is_file() else Path(f).rglob("*")) if s.is_file()))
    NOTES.append(f"{human(tot)} → {human(dst.stat().st_size)}")
    return [dst]


@action("📦 Extraire l'archive", ARC)
def extract_archive(files):
    outs = []
    for f in files:
        name = re.sub(r"\.(tar\.gz|tar\.bz2|tar\.xz|tgz|zip|tar|gz|bz2|xz)$", "", Path(f).name, flags=re.I)
        d = unique(Path(f).with_name(name))
        if zipfile.is_zipfile(f):
            d.mkdir()
            with zipfile.ZipFile(f) as z:
                for info in z.infolist():
                    if info.filename.startswith("__MACOSX/"):
                        continue
                    z.extract(info, d)
        elif tarfile.is_tarfile(f):
            d.mkdir()
            with tarfile.open(f) as t:
                t.extractall(d, filter="data")
        elif ext(f) == "gz":
            d = unique(Path(f).with_name(name))
            with gzip.open(f, "rb") as src, open(d, "wb") as o:
                shutil.copyfileobj(src, o)
        else:
            raise RuntimeError(f"Format d'archive non pris en charge : {Path(f).name}")
        outs.append(d)
    return outs


@action("🏷️ Renommer en lot (numérotation)", None, minf=2)
def batch_rename(files):
    base = ask(f"Renommer {len(files)} éléments.\nNom de base (ils deviendront « Nom_01 », « Nom_02 »…) :", "Fichier")
    if not base:
        raise Cancel()
    w = max(2, len(str(len(files))))
    preview = "\n".join(f"{Path(f).name}  →  {base}_{i:0{w}d}{Path(f).suffix}" for i, f in enumerate(files[:6], 1))
    if buttons(f"Aperçu :\n\n{preview}{chr(10) + '…' if len(files) > 6 else ''}", ["Annuler", "Renommer"]) != "Renommer":
        raise Cancel()
    outs = []
    tmp = []
    for i, f in enumerate(files, 1):  # passage par des noms temporaires pour éviter les collisions
        t = Path(f).with_name(f".ff_tmp_{i}_{os.getpid()}{Path(f).suffix}")
        os.rename(f, t)
        tmp.append(t)
    for i, t in enumerate(tmp, 1):
        dst = unique(t.with_name(f"{base}_{i:0{w}d}{t.suffix}"))
        os.rename(t, dst)
        outs.append(dst)
    return outs

# ───────────────────────── MENU PRINCIPAL ─────────────────────────

def finish(outs):
    outs = [Path(o) for o in outs if o]
    if not outs and not NOTES:
        return
    msg = f"Terminé ✓  {len(outs)} élément(s) créé(s)." if outs else "Terminé."
    if outs:
        msg += "\n\n" + "\n".join("• " + o.name for o in outs[:8]) + (f"\n… et {len(outs) - 8} autres" if len(outs) > 8 else "")
    if NOTES:
        msg += "\n\n" + "\n".join(NOTES)
    if not outs:
        buttons(msg, ["OK"])
        return
    btns = ["OK", "Afficher dans le Finder"] + (["Ouvrir"] if len(outs) == 1 else [])
    b = buttons(msg, btns, default=btns[-1])
    if b == "Ouvrir":
        open_path(outs[0])
    elif b.startswith("Afficher"):
        reveal(outs[0])


def available(files):
    res = []
    for a in ACTIONS:
        if len(files) < a["minf"]:
            continue
        if a["types"] is not None and not all(Path(f).is_file() and ext(f) in a["types"] for f in files):
            continue
        res.append(a)
    return res


def main(argv):
    files = sorted({os.path.abspath(a) for a in argv if os.path.exists(a)}, key=natural)
    if not files:
        alert("Aucun fichier", "Sélectionne un ou plusieurs fichiers dans le Finder, puis clic droit → Actions rapides → FileForge.")
        return
    acts = available(files)
    names = ", ".join(Path(f).name for f in files[:3]) + (f" (+{len(files) - 3})" if len(files) > 3 else "")
    label = choose([a["label"] for a in acts], f"{len(files)} élément(s) : {names}\n\nQue veux-tu faire ?", ok="Continuer")
    act = next(a for a in acts if a["label"] == label)
    log(f"{act['label']} ← {files}")
    outs = act["fn"](files) or []
    log(f"  → {[str(o) for o in outs]}")
    finish(outs)


def tk_entry(a):
    """Fenêtres visuelles lancées par la roue dans un processus séparé."""
    if IS_WIN:
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
        base = Path(sys.base_prefix)
        for var, pat in (("TCL_LIBRARY", "tcl8*"), ("TK_LIBRARY", "tk8*")):
            if var not in os.environ:
                hit = next((d for d in [*base.glob("tcl/" + pat), *base.glob("lib/" + pat)] if (d / "init.tcl").exists() or d.name.startswith("tk")), None)
                if hit:
                    os.environ[var] = str(hit)
    try:
        if a["fn"] == "regions":
            if a["kind"] == "pdf":
                import pymupdf
                src = pymupdf.open(a["path"])
            else:
                src = load_img(a["path"])
            r = pick_regions(src, a["kind"], a["title"], a["hint"], a.get("multi", False), a.get("start_page", 0))
            print(json.dumps({"ok": {str(k): [list(x) for x in v] for k, v in r.items()}}))
        else:
            print(json.dumps({"ok": str(draw_signature())}))
    except Cancel:
        print(json.dumps({"cancel": True}))
    except Exception as e:
        log(traceback.format_exc())
        print(json.dumps({"error": str(e)}))


if __name__ == "__main__" and len(sys.argv) > 2 and sys.argv[1] == "--tk":
    tk_entry(json.loads(sys.argv[2]))
    sys.exit(0)

if __name__ == "__main__":
    try:
        main(sys.argv[1:])
    except Cancel:
        pass
    except Exception as e:
        log(traceback.format_exc())
        alert("Oups, l'action a échoué", f"{e}\n\nDétails : ~/Library/Logs/FileForge.log")
